"""
exporter.py - Stage 6: write assignments back to the workbook.

Writes PCI / RSI / Mod4 into the Planning sheet (matched by site_sector_ID - never by row position, 
since row order isn't guaranteed to match assignments' order) and adds a KPI_Report sheet built straight
from validator.py's report dict. Doesn't touch any other existing sheet or column.
"""
from pathlib import Path
import openpyxl

# -- Internal helpers ---------------------------------------------------------------
def _write_planning_sheet(ws, assignments: dict[str, dict]) -> int:
    """
    Match every data row by its site_sector_ID cell and overwrite that row's PCI/RSI/Mod4 cells in place. 
    Returns how many rows were written.
    """
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))] # get the header row values
    col = {name: i for i, name in enumerate(header)}  # 0-indexed into row tuple - maps column name to its index in the row tuple
    sid_idx, pci_idx, rsi_idx, mod4_idx = col["site_sector_ID"], col["PCI"], col["RSI"], col["Mod4"] # get the column indices

    written = 0
    # iterate over all rows in the sheet starting from row 2 (skip header)
    for row in ws.iter_rows(min_row=2):
        sid = row[sid_idx].value # get the site_sector_ID value from the row
        if sid is None:
            continue # skip empty rows
        a = assignments.get(str(sid)) # get the assignment for this site_sector_ID
        if a is None:
            continue # skip rows that don't have an assignment
        # write the PCI, RSI, and Mod4 values into the corresponding cells in the row
        row[pci_idx].value = a["pci"]
        row[rsi_idx].value = a["rsi"]
        row[mod4_idx].value = a["mod4"]
        written += 1
    return written # return the number of rows written

def _write_kpi_sheet(wb, report):
    if "KPI_Report" in wb.sheetnames:
        del wb["KPI_Report"]
    ws = wb.create_sheet("KPI_Report")

    ws.append(["Metric", "Value"])
    ws.append(["Overall result", "PASS" if report["pass"] else "FAIL"])
    ws.append([])

    ws.append(["Hard rules (must be 0)", ""])
    for k, v in report["hard"].items():
        ws.append([k, v])
    ws.append([])

    ws.append(["Soft conflicts (informational)", ""])
    for k, v in report["soft"].items():
        ws.append([k, v])
    ws.append([])

    ws.append(["Full KPI breakdown (all granularities)", ""])
    for k, v in report["counts"].items():
        ws.append([k, v])
    ws.append([])

    ws.append(["Coverage", ""])
    for k, v in report["coverage"].items():
        ws.append([k, v])

    # Auto-width columns
    for col_cells in ws.columns:
        width = max(len(str(c.value)) for c in col_cells if c.value is not None) + 2
        ws.column_dimensions[col_cells[0].column_letter].width = max(width, 12)

# -- Public API ----------------------------------------------------------------------
def export_results(
    src_path: str | Path,
    out_path: str | Path,
    assignments: dict[str, dict],
    report: dict,
    sheet_name: str = "Planning",
) -> tuple[str, int]:
    """
    Reads `src_path`, writes `assignments` into its Planning sheet, adds a
    KPI_Report sheet from `report` (see validator.validate), and saves to `out_path`. 
    Returns (out_path, rows_written).
    """
    # load the workbook and check if the specified sheet exists
    wb = openpyxl.load_workbook(src_path)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"sheet '{sheet_name}' not found; available: {wb.sheetnames}")

    written = _write_planning_sheet(wb[sheet_name], assignments)
    _write_kpi_sheet(wb, report)

    # ensure the output directory exists and save the workbook
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return str(out_path), written # return the output path and the number of rows written

# -- Command-line test -------------------------------------------------------
if __name__ == "__main__":
    import sys
    from bulk_plan import bulk_plan
    from validator import validate

    src = sys.argv[1] if len(sys.argv) > 1 else "data/NR5G_PCI_RSI_Planning.xlsx"
    out = sys.argv[2] if len(sys.argv) > 2 else "output/NR5G_PCI_RSI_Planning_final(mod_final1).xlsx"

    sectors, graph, assignments = bulk_plan(src)
    report = validate(sectors, graph, assignments)
    out_path, written = export_results(src, out, assignments, report)
    print(f"Wrote {written} sector rows + KPI_Report sheet to {out_path}")
    print("PASS" if report["pass"] else "FAIL - hard rule violated")